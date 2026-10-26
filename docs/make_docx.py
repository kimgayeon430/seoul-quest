"""docs/report.md → docs/report.docx (A4, 한글 글꼴, 표 테두리, 그림 크기·쪽 나눔 정리).

- 그림은 쪽 안에 들어가도록 폭 6.0in·높이 7.0in 이내로 줄이고, 캡션과 같은 쪽에 붙인다.
- 장(##)마다 새 쪽에서 시작하고(초록·Abstract 제외), 표는 행이 쪽 경계에서 갈라지지 않게 한다.

    pip install pypandoc_binary python-docx
    python docs/make_docx.py
"""

import re
import subprocess
import tempfile
from pathlib import Path

import pypandoc
from docx import Document
from docx.oxml.ns import qn
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

DOCS = Path(__file__).resolve().parent
KO_FONT, MONO_FONT = "Malgun Gothic", "Consolas"
MAX_W, MAX_H = Inches(6.0), Inches(7.0)
NO_BREAK_BEFORE = ("초록", "Abstract")


def main() -> None:
    pandoc = pypandoc.get_pandoc_path()
    tmp = Path(tempfile.mkdtemp())
    subprocess.run([pandoc, "-o", tmp / "ref.docx", "--print-default-data-file", "reference.docx"], check=True)
    ref = Document(tmp / "ref.docx")
    for st in ref.styles:
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rf)
        mono = st.name in ("Source Code", "Verbatim Char")
        for a in ("w:ascii", "w:hAnsi"):
            rf.set(qn(a), MONO_FONT if mono else KO_FONT)
        rf.set(qn("w:eastAsia"), KO_FONT)
    ref.styles["Normal"].font.size = Pt(10.5)
    ref.save(tmp / "ref_ko.docx")

    # "![짧은 alt](경로)" 다음 줄의 기울임 캡션을 그림 캡션(alt)으로 옮긴다.
    md = re.sub(r"!\[[^\]]*\]\(([^)]+)\)\n\n\*(그림 \d+\.[^\n]*)\*\n", r"![\2](\1)\n",
                (DOCS / "report.md").read_text())
    (tmp / "report.md").write_text(md)
    subprocess.run([pandoc, tmp / "report.md", "-f", "markdown-tex_math_dollars-raw_tex", "-o", tmp / "raw.docx",
                    "--reference-doc", tmp / "ref_ko.docx", "--resource-path", DOCS], check=True)

    d = Document(tmp / "raw.docx")
    for sec in d.sections:
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.5)
        sec.top_margin = sec.bottom_margin = Cm(2.5)

    for t in d.tables:
        borders = t._tbl.tblPr.makeelement(qn("w:tblBorders"), {})
        for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
            borders.append(borders.makeelement(qn("w:" + e), {qn("w:val"): "single", qn("w:sz"): "4",
                                                               qn("w:space"): "0", qn("w:color"): "999999"}))
        t._tbl.tblPr.append(borders)
        for i, row in enumerate(t.rows):
            trpr = row._tr.get_or_add_trPr()
            trpr.append(trpr.makeelement(qn("w:cantSplit"), {}))
            if i == 0:
                trpr.append(trpr.makeelement(qn("w:tblHeader"), {}))
            last = i == len(t.rows) - 1
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = not last  # 짧은 표는 한 쪽에 모은다
                    for r in p.runs:
                        r.font.size = Pt(9)

    for s in d.inline_shapes:
        ratio = min(1.0, MAX_W / s.width, MAX_H / s.height)
        if ratio < 1.0:
            s.width, s.height = int(s.width * ratio), int(s.height * ratio)

    for p in d.paragraphs:
        style = p.style.name
        if p._p.xpath(".//w:drawing"):  # 그림 문단: 가운데, 캡션과 같은 쪽
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.keep_together = True
        elif style in ("Image Caption", "Caption"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(12)
        elif style == "Source Code":
            p.paragraph_format.keep_together = True
        elif style.startswith("Heading"):
            p.paragraph_format.keep_with_next = True
            if style == "Heading 2" and p.text.strip() not in NO_BREAK_BEFORE:
                p.paragraph_format.page_break_before = True
        elif style in ("Table Caption",):
            p.paragraph_format.keep_with_next = True

    d.core_properties.title = "온디바이스 다중 신호 사진 인증과 하이브리드 추천을 적용한 여행 미션 게이미피케이션 앱"
    d.core_properties.author = "김가연, 강규린"
    d.save(DOCS / "report.docx")
    print("저장:", DOCS / "report.docx")


if __name__ == "__main__":
    main()
